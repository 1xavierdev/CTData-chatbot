"""Turn a chat answer (text, sources and optional chart) into a PowerPoint deck or a PDF.

    doc = {"title": "What is the chronic absenteeism rate?", "answer": "<markdown>",
           "sources": [{"title": ..., "url": ...}], "chart": <charts.py spec> | None,
           "chart_png": "data:image/png;base64,..." | None}   # the chart as drawn in the browser

PowerPoint charts are native (editable in PowerPoint); the PDF embeds the browser's
chart image and always adds the numbers as a table.
"""

from __future__ import annotations

import base64
import io
import re
from datetime import date
from pathlib import Path

from fpdf import FPDF
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.util import Inches, Pt

NAVY = RGBColor(0x0B, 0x2A, 0x4A)
GRAY = RGBColor(0x5B, 0x67, 0x75)
CHART_TYPES = {"line": XL_CHART_TYPE.LINE_MARKERS, "bar": XL_CHART_TYPE.COLUMN_CLUSTERED,
               "doughnut": XL_CHART_TYPE.DOUGHNUT}
BLOCKS_PER_SLIDE = 6
# A Unicode font makes curly quotes, dashes and bullets print as-is; without one, text is simplified.
FONT_DIR = Path("/System/Library/Fonts/Supplemental")


def blocks(markdown: str) -> list[tuple[str, str]]:
    """Split the model's Markdown into ("p" | "li", text) blocks."""
    out = []
    for line in markdown.splitlines():
        line = line.strip()
        if not line:
            continue
        if bullet := re.match(r"^(?:[-*•]|\d+\.)\s+(.*)$", line):
            out.append(("li", bullet.group(1)))
        else:
            out.append(("p", re.sub(r"^#+\s*", "", line)))
    return out


def plain(text: str) -> str:
    return re.sub(r"\*\*(.+?)\*\*", r"\1", text).replace("`", "")


def subtitle() -> str:
    return f"CTData Assistant · {date.today():%B %-d, %Y}"


# ---------------------------------------------------------------- PowerPoint

def add_runs(paragraph, text: str, size: int) -> None:
    """Add text to a paragraph, turning **bold** into bold runs."""
    for i, part in enumerate(re.split(r"\*\*(.+?)\*\*", text.replace("`", ""))):
        if part:
            run = paragraph.add_run()
            run.text = part
            run.font.size = Pt(size)
            run.font.bold = i % 2 == 1


def place(shape, left: float, top: float, width: float, height: float) -> None:
    shape.left, shape.top, shape.width, shape.height = Inches(left), Inches(top), Inches(width), Inches(height)


def new_slide(prs, layout: int, title: str, size: int = 30):
    """A slide from the default template, re-laid out for 16:9 (its placeholders are sized for 4:3)."""
    slide = prs.slides.add_slide(prs.slide_layouts[layout])
    bar = slide.shapes.add_shape(1, 0, 0, prs.slide_width, Inches(0.18))  # 1 = rectangle
    bar.fill.solid()
    bar.fill.fore_color.rgb = NAVY
    bar.line.fill.background()
    slide.shapes.title.text = title
    place(slide.shapes.title, 0.7, 0.45, 11.9, 1.0)
    for paragraph in slide.shapes.title.text_frame.paragraphs:
        paragraph.alignment = 1  # left
        for run in paragraph.runs:
            run.font.size, run.font.bold, run.font.color.rgb = Pt(size), True, NAVY
    return slide


def build_pptx(doc: dict) -> bytes:
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)

    slide = new_slide(prs, 0, doc.get("title") or "CTData Assistant", 40)
    place(slide.shapes.title, 0.9, 2.3, 11.5, 2.0)
    slide.placeholders[1].text = subtitle()
    place(slide.placeholders[1], 0.9, 4.4, 11.5, 0.8)
    slide.placeholders[1].text_frame.paragraphs[0].alignment = 1
    slide.placeholders[1].text_frame.paragraphs[0].runs[0].font.color.rgb = GRAY

    parts = blocks(doc.get("answer", ""))
    for start in range(0, len(parts), BLOCKS_PER_SLIDE):
        slide = new_slide(prs, 1, "Answer" if start == 0 else "Answer (continued)")
        body = slide.placeholders[1]
        place(body, 0.7, 1.6, 11.9, 5.4)
        frame = body.text_frame
        frame.word_wrap = True
        for i, (kind, text) in enumerate(parts[start:start + BLOCKS_PER_SLIDE]):
            paragraph = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
            paragraph.level = 1 if kind == "li" else 0
            add_runs(paragraph, text, 18 if kind == "p" else 16)

    if chart := doc.get("chart"):
        slide = new_slide(prs, 5, chart.get("title") or "Chart", 24)
        data = CategoryChartData()
        data.categories = chart["labels"]
        for s in chart["series"]:
            data.add_series(s["name"], s["values"])
        kind = chart.get("type", "bar")
        graphic = slide.shapes.add_chart(CHART_TYPES.get(kind, XL_CHART_TYPE.COLUMN_CLUSTERED),
                                         Inches(0.8), Inches(1.5), Inches(11.7), Inches(5.4), data).chart
        graphic.has_legend = len(chart["series"]) > 1 or kind == "doughnut"
        if graphic.has_legend:
            graphic.legend.position = XL_LEGEND_POSITION.BOTTOM
            graphic.legend.include_in_layout = False
        if kind == "doughnut":
            plot = graphic.plots[0]
            plot.has_data_labels = True
            plot.data_labels.number_format = '0.0"%"' if chart.get("unit") == "%" else "#,##0"
            plot.data_labels.number_format_is_linked = False
        else:
            graphic.value_axis.has_major_gridlines = True
            graphic.value_axis.tick_labels.font.size = Pt(12)
            graphic.category_axis.tick_labels.font.size = Pt(12)

    if sources := doc.get("sources"):
        slide = new_slide(prs, 1, "Sources")
        place(slide.placeholders[1], 0.7, 1.6, 11.9, 5.4)
        frame = slide.placeholders[1].text_frame
        for i, source in enumerate(sources):
            paragraph = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
            run = paragraph.add_run()
            run.text = source.get("title") or source["url"]
            run.hyperlink.address = source["url"]
            run.font.size = Pt(16)
        note = frame.add_paragraph()
        add_runs(note, f"Generated by the {subtitle()} from ctdata.org content.", 12)
        note.runs[0].font.color.rgb = GRAY

    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()


# ---------------------------------------------------------------- PDF

class AnswerPDF(FPDF):
    def setup_fonts(self) -> None:
        regular, bold = FONT_DIR / "Arial.ttf", FONT_DIR / "Arial Bold.ttf"
        if regular.exists() and bold.exists():
            self.add_font("Body", "", str(regular))
            self.add_font("Body", "B", str(bold))
            self.font_name, self.unicode = "Body", True
        else:
            self.font_name, self.unicode = "Helvetica", False

    def text(self, value: str) -> str:  # noqa: A003 - fpdf's own text() isn't used here
        if self.unicode:
            return value
        table = str.maketrans({"‘": "'", "’": "'", "“": '"', "”": '"', "–": "-", "—": "-", "…": "...", "•": "-", "·": "-"})
        return value.translate(table).encode("latin-1", "replace").decode("latin-1")

    def footer(self) -> None:
        self.set_y(-12)
        self.set_font(self.font_name, "", 8)
        self.set_text_color(120, 128, 138)
        self.cell(0, 6, self.text(f"{subtitle()}  |  page {self.page_no()}"), align="C")


def build_pdf(doc: dict) -> bytes:
    pdf = AnswerPDF(format="A4")
    pdf.setup_fonts()
    pdf.set_margins(18, 18, 18)
    pdf.set_auto_page_break(True, margin=18)
    pdf.add_page()
    f = pdf.font_name
    width = pdf.epw

    pdf.set_font(f, "", 9)
    pdf.set_text_color(91, 103, 117)
    pdf.cell(width, 5, pdf.text(subtitle()), new_x="LMARGIN", new_y="NEXT")
    pdf.ln(2)
    pdf.set_font(f, "B", 16)
    pdf.set_text_color(11, 42, 74)
    pdf.multi_cell(width, 8, pdf.text(doc.get("title") or "CTData Assistant"), align="L", new_x="LMARGIN", new_y="NEXT")
    pdf.ln(3)

    pdf.set_text_color(30, 30, 30)
    pdf.set_font(f, "", 11)
    for kind, text in blocks(doc.get("answer", "")):
        text = pdf.text(text.replace("--", "-").replace("__", "_"))  # keep fpdf's markdown to **bold** only
        if kind == "li":
            pdf.set_x(pdf.l_margin + 4)
            pdf.multi_cell(width - 4, 6, "•  " + text if pdf.unicode else "-  " + text,
                           markdown=True, align="L", new_x="LMARGIN", new_y="NEXT")
        else:
            pdf.multi_cell(width, 6, text, markdown=True, align="L", new_x="LMARGIN", new_y="NEXT")
        pdf.ln(1.5)

    if chart := doc.get("chart"):
        png = doc.get("chart_png") or ""
        # Keep the chart's heading on the same page as the chart (or its table).
        needed = 14 + (width * 0.5 if png.startswith("data:image/png;base64,") else 30)
        if pdf.get_y() + needed > pdf.page_break_trigger:
            pdf.add_page()
        pdf.ln(3)
        pdf.set_font(f, "B", 12)
        pdf.set_text_color(11, 42, 74)
        pdf.multi_cell(width, 7, pdf.text(chart.get("title") or "Chart"), align="L", new_x="LMARGIN", new_y="NEXT")
        if png.startswith("data:image/png;base64,"):
            pdf.image(io.BytesIO(base64.b64decode(png.split(",", 1)[1])), w=width)
            pdf.ln(3)
        # The numbers as a table: one row per label, one column per series.
        pdf.set_font(f, "", 9)
        pdf.set_text_color(30, 30, 30)
        unit = f" ({chart['unit']})" if chart.get("unit") else ""
        values = [v for s in chart["series"] for v in s["values"] if v is not None]
        decimals = 1 if any(v != int(v) for v in values) else 0  # 51.0 next to 49.1, not 51
        with pdf.table(first_row_as_headings=True, text_align="LEFT", line_height=6) as table:
            table.row([pdf.text(t) for t in ["", *(s["name"] + unit for s in chart["series"])]])
            for i, label in enumerate(chart["labels"]):
                table.row([pdf.text(str(label)), *("" if s["values"][i] is None else f"{s['values'][i]:,.{decimals}f}"
                                                   for s in chart["series"])])

    if sources := doc.get("sources"):
        pdf.ln(5)
        pdf.set_font(f, "B", 12)
        pdf.set_text_color(11, 42, 74)
        pdf.cell(width, 7, "Sources", new_x="LMARGIN", new_y="NEXT")
        pdf.set_font(f, "", 10)
        pdf.set_text_color(21, 101, 192)
        for source in sources:
            pdf.multi_cell(width, 6, pdf.text(source.get("title") or source["url"]), link=source["url"],
                           new_x="LMARGIN", new_y="NEXT")
    return bytes(pdf.output())
