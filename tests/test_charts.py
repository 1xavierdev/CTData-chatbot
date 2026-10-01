import io
import unittest
from unittest import mock

from pptx import Presentation

import charts
import export
import rag
import server
from kb import KnowledgeBase, build_chunks

TABLE_URL = "https://public-edsight.ct.gov/performance/smarter-balanced"
TABLES = {TABLE_URL: {"title": "EdSight: Smarter Balanced (statewide)", "records": [
    {"name": "State of Connecticut, ELA", "year": "2024-25", "measure": "Total Number with Scored Tests", "value": "217,969"},
    {"name": "State of Connecticut, ELA", "year": "2024-25", "measure": "Percentage Level 3 or 4 %", "value": "50.3"},
    {"name": "State of Connecticut, ELA", "year": "2025-26", "measure": "Percentage Level 3 or 4 %", "value": "51.0"},
    {"name": "State of Connecticut, Math", "year": "2024-25", "measure": "Percentage Level 3 or 4 %", "value": "45.9"},
    {"name": "State of Connecticut, Math", "year": "2025-26", "measure": "Percentage Level 3 or 4 %", "value": "46.7"},
]}}
PASSAGES = [{"id": f"{TABLE_URL}#0", "url": TABLE_URL, "title": "EdSight: Smarter Balanced (statewide)", "text": "..."}]
CHART = {"type": "line", "title": "Scores", "unit": "%", "labels": ["2024-25", "2025-26"],
         "series": [{"name": "ELA", "values": [50.3, 51.0]}, {"name": "Math", "values": [45.9, 46.7]}]}
DOC = {"title": "Compare ELA and math", "answer": "Both **improved**:\n\n* ELA: 51.0%\n* Math: 46.7%",
       "sources": [{"title": "EdSight", "url": TABLE_URL}], "chart": CHART}


class ChartTests(unittest.TestCase):
    def setUp(self):
        patcher = mock.patch("charts.tables", return_value=TABLES)
        patcher.start()
        self.addCleanup(patcher.stop)

    def test_intent(self):
        self.assertEqual(charts.chart_intent("Draw a graph of evictions"), "explicit")
        self.assertEqual(charts.chart_intent("Compare ELA vs math"), "compare")
        self.assertIsNone(charts.chart_intent("What is PSEO?"))

    def test_edsight_comparison_uses_structured_percentages(self):
        chart = charts.build_chart("Compare ELA and math results", PASSAGES, use_model=False)
        self.assertEqual(chart["labels"], ["2024-25", "2025-26"])
        self.assertEqual([s["name"] for s in chart["series"]], ["ELA", "Math"])
        self.assertEqual(chart["series"][1]["values"], [45.9, 46.7])
        self.assertEqual(chart["unit"], "%")

    def test_named_row_is_charted_alone(self):
        chart = charts.build_chart("What percent met the math standard?", PASSAGES, use_model=False)
        self.assertEqual([s["name"] for s in chart["series"]], ["Math"])

    def test_no_chart_without_intent_or_table(self):
        other = [{**PASSAGES[0], "url": "https://www.ctdata.org/blog/x"}]
        self.assertIsNone(charts.build_chart("Compare evictions", other, use_model=False))
        self.assertIsNone(charts.build_chart("Who runs EdSight?", PASSAGES, use_model=False))

    def test_pick_type(self):
        self.assertEqual(charts.pick_type(["2019", "2020", "2021"], [{"values": [1, 2, 3]}], ""), "line")
        self.assertEqual(charts.pick_type(["A", "B"], [{"values": [60, 40]}], "%"), "doughnut")
        self.assertEqual(charts.pick_type(["A", "B"], [{"values": [6, 4]}], "cases"), "bar")


class ModelChartValidationTests(unittest.TestCase):
    CONTEXT = "Filings rose from 13,847 in 2010 to 44,146 in 2024."

    def test_values_must_appear_in_context(self):
        good = {"title": "t", "unit": "", "labels": ["2010", "2024"], "series": [{"name": "n", "values": [13847, 44146]}]}
        made_up = {**good, "series": [{"name": "n", "values": [13847, 50000]}]}
        self.assertEqual(charts.validate(good, self.CONTEXT)["series"][0]["values"], [13847, 44146])
        self.assertIsNone(charts.validate(made_up, self.CONTEXT))

    def test_flat_filler_and_mismatched_lengths_rejected(self):
        flat = {"title": "t", "unit": "", "labels": ["2010", "2024"], "series": [{"name": "n", "values": [2010, 2010]}]}
        short = {"title": "t", "unit": "", "labels": ["2010", "2024", "x"], "series": [{"name": "n", "values": [13847, 44146]}]}
        self.assertIsNone(charts.validate(flat, self.CONTEXT))
        self.assertIsNone(charts.validate(short, self.CONTEXT))

    def test_numbers_must_be_stated_with_their_label(self):
        context = ("In 2021 and 2022, males accounted for 52% and 51% of deaths. "
                   "Hartford (16,129), Bridgeport (12,243) and New Haven (10,006) led filings.")
        wrong_years = {"title": "t", "unit": "%", "labels": ["2019", "2020"], "series": [{"name": "Males", "values": [52, 51]}]}
        right_years = {**wrong_years, "labels": ["2021", "2022"]}
        towns = {"title": "t", "unit": "", "labels": ["Hartford", "New Haven"], "series": [{"name": "Filings", "values": [16129, 10006]}]}
        swapped = {**towns, "labels": ["New Haven", "Hartford"]}
        self.assertIsNone(charts.validate(wrong_years, context))
        self.assertIsNotNone(charts.validate(right_years, context))
        self.assertEqual(charts.validate(towns, context)["type"], "bar")
        self.assertIsNone(charts.validate(swapped, context))

    def test_counts_are_not_percentages(self):
        spec = {"title": "t", "unit": "%", "labels": ["2010", "2024"], "series": [{"name": "n", "values": [13847, 44146]}]}
        self.assertEqual(charts.validate(spec, self.CONTEXT)["unit"], "")


class ExportTests(unittest.TestCase):
    def test_pptx_has_answer_chart_and_sources(self):
        prs = Presentation(io.BytesIO(export.build_pptx(DOC)))
        titles = [s.shapes.title.text for s in prs.slides]
        self.assertEqual(titles, ["Compare ELA and math", "Answer", "Scores", "Sources"])
        self.assertTrue(any(shape.has_chart for shape in prs.slides[2].shapes))

    def test_pdf_is_built(self):
        self.assertTrue(export.build_pdf(DOC).startswith(b"%PDF"))

    def test_export_endpoint(self):
        client = server.app.test_client()
        resp = client.post("/api/export", json={**DOC, "format": "pdf"})
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp.mimetype, "application/pdf")
        self.assertIn("compare-ela-and-math.pdf", resp.headers["Content-Disposition"])
        self.assertEqual(client.post("/api/export", json={**DOC, "format": "docx"}).status_code, 400)


class ExportIntentTests(unittest.TestCase):
    def setUp(self):
        self.kb = KnowledgeBase(build_chunks([{"url": "https://www.ctdata.org/education", "title": "Education",
                                               "text": "Chronic Absenteeism statewide", "links": []}]))

    def test_export_previous_answer(self):
        result = rag.answer(self.kb, "make a ppt of this", {})
        self.assertEqual((result["mode"], result["export"], result["export_previous"]), ("export", "pptx", True))

    def test_export_new_answer(self):
        with mock.patch("lmstudio.chat", return_value="Chronic absenteeism is 16.4%."):
            result = rag.answer(self.kb, "create a pdf about chronic absenteeism", {})
        self.assertEqual((result["mode"], result["export"]), ("llm", "pdf"))
        self.assertNotIn("export_previous", result)


if __name__ == "__main__":
    unittest.main()
